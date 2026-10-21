"""Build the Midterm slide deck and its Vietnamese speaker script from one source (SLIDES below).

Slide text is English (matches the report and the brief); the speaker notes are Vietnamese.
Every number comes from results/tables/*.csv; figures come from results/figures/.

Outputs (next to the repo, in the workspace folder):
    outputs/midterm_slides.pptx
    research/presentation-script.md

Usage:
    python scripts/build_presentation.py
"""
import sys
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

REPO = Path(__file__).resolve().parents[1]
OUT_DIR = REPO.parent / "outputs"
SCRIPT_MD = REPO.parent / "research" / "presentation-script.md"
FIG = REPO / "results" / "figures"
TAB = REPO / "results" / "tables"

ms = pd.read_csv(TAB / "midterm_ms_baselines.csv").set_index("model")
bl = pd.read_csv(TAB / "midterm_baselines.csv")
e3 = pd.read_csv(TAB / "midterm_ms_e3_crnn_vs_cnn_by_duration.csv")
e2 = pd.read_csv(TAB / "midterm_ms_e2_confusion_pairs.csv").head(5)
par = pd.read_csv(TAB / "rq4_pareto.csv").set_index("model")


def f(m):
    return f"{ms.loc[m, 'test_macro_f1_mean']:.3f} ± {ms.loc[m, 'test_macro_f1_std']:.3f}"


SLIDES = [
    {
        "title": "DataSEC: Environmental Sound Classification under Noise",
        "bullets": ["ATI 2026 · Track 3 · Midterm", "22-class log-Mel classification: dataset, baselines, error analysis"],
        "notes": "Chào thầy/cô và các bạn. Nhóm em làm track 3: phân loại âm thanh môi trường 22 lớp trên bộ DataSEC. "
                 "Hôm nay em trình bày phần giữa kỳ: hiểu dữ liệu, pipeline, ba mô hình nền, và phân tích lỗi; "
                 "cuối cùng là câu hỏi nghiên cứu cho phần Final.",
    },
    {
        "title": "Dataset: 5,048 clips, 22 classes, heavily imbalanced",
        "bullets": ["Zenodo DOI 10.5281/zenodo.17033970 (pinned), mono 44.1 kHz, 23.7 h",
                    "Brief says 4,292 clips; the pinned archive has 5,048 (we report the difference)",
                    "Voices 1,900 vs Cat fights 50 (38:1); Voices + Music = 57.5 % of clips",
                    "Predict-all-Voices: accuracy 0.376 but macro-F1 0.025, so macro-F1 is the main metric"],
        "image": "eda_class_counts.png",
        "notes": "Bộ dữ liệu có 5.048 file, nhiều hơn con số 4.292 trong đề vì bản Zenodo ghim DOI là bản mới; nhóm đã ghi rõ chênh lệch này. "
                 "Mất cân bằng rất nặng: lớp lớn nhất gấp 38 lần lớp nhỏ nhất. Chỉ cần đoán toàn bộ là Voices đã được accuracy 0,376 "
                 "nhưng macro-F1 chỉ 0,025. Vì vậy chúng em dùng macro-F1 làm thước đo chính và dùng loss có trọng số theo lớp.",
    },
    {
        "title": "Durations vary from 0.2 s to 8 min, and duration leaks the class",
        "bullets": ["With a 5 s window: 38.3 % of clips are zero-padded, 60.0 % are cropped",
                    "Duration alone (gradient boosting on log-duration) gives macro-F1 0.272 vs 0.045 chance",
                    "Curation artefacts: 88 % of Voices and 72 % of Music are cut off at 11.025 kHz; 93 % of Wind turbine clips are exactly 10 s"],
        "image": "eda_duration_box.png",
        "notes": "Thời lượng rất khác nhau giữa các lớp, từ 0,2 giây đến 8 phút. Chỉ dùng thời lượng để đoán đã cho macro-F1 0,272, cao hơn ngẫu nhiên nhiều, "
                 "nghĩa là mô hình có thể học độ dài phần đệm thay vì học âm thanh. Ngoài ra Voices và Music bị cắt cứng ở 11 kHz, nên ta resample về 22,05 kHz để xoá dấu vết này. "
                 "Đây là các rủi ro shortcut mà chúng em nhớ khi đọc kết quả.",
    },
    {
        "title": "Waveform vs linear spectrogram vs log-Mel",
        "bullets": ["Waveform: pressure over time, no frequency axis, phase-sensitive",
                    "Linear spectrogram |STFT|²: 513 uniform bins (n_fft 1024, hop 512)",
                    "log-Mel: 64 triangular Mel filters, then 10·log10; gain becomes an additive constant",
                    "Network input: 1 × 64 × 216"],
        "image": "eda_representations.png",
        "notes": "Waveform cho biết khi nào to nhỏ nhưng không đọc được tần số. Spectrogram tuyến tính có 513 bin chia đều, phần lớn là tần số cao ít thông tin. "
                 "log-Mel gom thành 64 dải theo thang Mel, dải thấp hẹp dải cao rộng giống tai người, rồi lấy log để nén dải động. "
                 "Một hệ quả quan trọng: đổi âm lượng chỉ là cộng một hằng số trong miền log, điều này liên quan đến thí nghiệm gain ở phần Final.",
    },
    {
        "title": "Which classes overlap spectrally?",
        "bullets": ["Cosine similarity of class-mean log-Mel spectra (level and common tilt removed)",
                    "Most similar: Jet–Propeller 0.992, Propeller–Train 0.988, Jet–Train 0.976",
                    "Two clusters: vehicles / engines, and bright sharp sounds (Glass, Workshop, Horn)",
                    "Spectral similarity is only part of the story: Glass–Workshop (0.982) is never confused"],
        "image": "eda_spectral_similarity.png",
        "notes": "Chúng em tính độ giống nhau của phổ trung bình giữa các lớp. Cặp giống nhất là Jet và Propeller, rồi Train. Chúng hình thành nhóm động cơ/xe cộ. "
                 "Tuy nhiên giống phổ chưa chắc bị nhầm: Glass breaking và Workshop giống 0,98 nhưng không bao giờ bị nhầm, vì tiếng thủy tinh là xung ngắn còn phổ trung bình đã mất thông tin thời gian.",
    },
    {
        "title": "Pipeline and split",
        "bullets": ["wav → resample 22.05 kHz → centre crop / zero-pad 5 s → STFT → 64 Mel → 10·log10 (80 dB floor) → train-set standardisation",
                    "Computed at runtime; only int16 audio is cached, no spectrogram files",
                    "Stratified split on 40 leaf classes, 70/15/15, fixed IDs in splits/; test = 758 clips",
                    "Checkpoints chosen on validation macro-F1; test used once. 2 test clips share audio with train (effect ≤ 0.0017 macro-F1)"],
        "notes": "Pipeline chạy lúc chạy, không lưu ảnh spectrogram: chỉ cache âm thanh int16. Split chia theo 40 lớp con nên lớp nhỏ nhất vẫn có 8 clip test. "
                 "Chọn checkpoint bằng macro-F1 trên tập validation, không bao giờ chỉnh theo test. Em cũng nói thẳng: có 2 clip test trùng một phần âm thanh với train; "
                 "loại chúng ra thì macro-F1 đổi không quá 0,0017 nên giữ nguyên split.",
    },
    {
        "title": "Three baseline models",
        "bullets": ["small_cnn: 4 conv blocks, 0.39 M params, 198 M MACs",
                    "mobilenetv3s: ImageNet-pretrained, 1→3 channel copy, 1.54 M, 18 M MACs",
                    "crnn: 4 conv blocks → BiGRU(128), 0.87 M, 487 M MACs",
                    "AdamW, cosine schedule, class-weighted cross-entropy, early stopping on val macro-F1; seeds 42, 43, 44"],
        "notes": "Ba mô hình theo đề: CNN nhỏ, MobileNetV3-Small pretrained trên ImageNet, và CRNN có GRU hai chiều. "
                 "Cùng một công thức huấn luyện: AdamW, cosine, loss trọng số theo lớp, dừng sớm theo macro-F1 validation. Chạy ba seed.",
    },
    {
        "title": "Baseline results: three seeds, no clear winner",
        "table": [["model", "test macro-F1", "test acc", "params", "MACs", "latency (ms)"],
                  ["small_cnn", f("small_cnn"), f"{ms.loc['small_cnn','test_acc_mean']:.3f}", "0.39 M", "198 M", f"{par.loc['small_cnn','latency_ms']:.2f}"],
                  ["mobilenetv3s", f("mobilenetv3s"), f"{ms.loc['mobilenetv3s','test_acc_mean']:.3f}", "1.54 M", "18 M", f"{par.loc['mobilenetv3s','latency_ms']:.2f}"],
                  ["crnn", f("crnn"), f"{ms.loc['crnn','test_acc_mean']:.3f}", "0.87 M", "487 M", f"{par.loc['crnn','latency_ms']:.2f}"]],
        "bullets": ["Mean ± std over seeds 42/43/44; seed spread (up to 0.027) is as large as the model gaps",
                    "The first single-seed Colab run ranked small_cnn 0.884 > mobilenet 0.853 > crnn 0.815; it did not replicate",
                    "Cause: that CRNN run early-stopped at epoch 20; local runs trained 40–50 epochs"],
        "notes": "Đây là kết quả chính. Với ba seed, ba mô hình gần như ngang nhau: 0,876, 0,901 và 0,875, độ lệch giữa seed lên tới 0,027 bằng khoảng cách giữa các mô hình. "
                 "Em muốn nói thẳng: lần chạy đầu trên Colab chỉ một seed cho thứ hạng khác, CRNN thấp nhất 0,815, nhưng lần đó CRNN dừng sớm ở epoch 20. "
                 "Khi chạy lại ba seed thì thứ hạng đó không lặp lại, nên nhóm không kết luận mô hình nào tốt hơn trên dữ liệu sạch. "
                 "Latency đo trên RTX 2060 SUPER, không so được với số T4 trước đó.",
    },
    {
        "title": "Per-class F1: best and worst classes",
        "bullets": ["Best (≥ 0.97): Voices, Chicken coop, Wind turbine, Music, Cicadas",
                    "Worst: Jet 0.71, Train 0.73, Vehicle idling 0.74, Propeller / Vacuum-fan 0.76",
                    "Hard classes are the same across seeds and models (std 0.06–0.09)",
                    "Caution: Voices, Music, Wind turbine also carry curation shortcuts"],
        "image": "test_per_class_f1.png",
        "notes": "Câu hỏi E1. Tốt nhất là Voices, Chicken coop, Wind turbine, Music, Cicadas với F1 từ 0,97. Tệ nhất là Jet, Train, Vehicle idling, Propeller và Vacuum. "
                 "Nhóm khó khá ổn định giữa các seed. Nhưng cần thận trọng: ba lớp điểm cao có dấu vết xử lý dữ liệu như cắt ở 11 kHz hay độ dài cố định, nên điểm cao chưa chắc là đã học đúng âm thanh. "
                 "Biểu đồ trên là seed 42.",
    },
    {
        "title": "Error analysis: engine-like sounds confuse each other",
        "table": [["class pair", "extra errors per run"]] +
                 [[f"{r.A} – {r.B}", f"{r.total_per_run:.1f}"] for r in e2.itertuples()],
        "bullets": ["Summed over 3 models × 3 seeds, divided by 9 runs",
                    "Nearly all are steady broadband machine / engine sounds",
                    "Music → Bells is the main non-engine confusion"],
        "image": "confusion_small_cnn.png",
        "notes": "Câu hỏi E2. Các cặp nhầm nhiều nhất đều là tiếng máy móc, động cơ liên tục: Jet với Propeller, Lawn mower với Vehicle idling, Vacuum với Workshop. "
                 "Ngoài nhóm này, lỗi đáng kể nhất là Music bị đoán thành Bells. Ma trận nhầm lẫn bên phải là small_cnn seed 42. "
                 "[CẦN BẠN LÀM] Phần nghe thử các clip bị nhầm chưa được điền; trước khi trình bày hãy nghe vài clip trong bảng E2 để nêu nguyên nhân thật.",
    },
    {
        "title": "Does the GRU help? CRNN − CNN by clip duration",
        "table": [["duration group", "n", "Δ macro-F1 mean ± std (3 seeds)"]] +
                 [[r.group, str(r.n), f"{r.delta_mean:+.3f} ± {r.delta_std:.3f}"] for r in e3.itertuples()],
        "bullets": ["No consistent difference in any group (signs change between seeds)",
                    "CRNN and CNN also differ in conv front end (487 M vs 198 M MACs), so this is not a clean test of recurrence",
                    "Controlled variant (small_cnn trunk + BiGRU) is a candidate for the Final"],
        "notes": "Câu hỏi E3. Chênh lệch CRNN trừ CNN theo độ dài clip nhỏ và đổi dấu giữa các seed, nên không có cải thiện đo được từ GRU. "
                 "Kết luận từ lần chạy một seed là CRNN kém hơn ở clip dài không còn đúng. Hơn nữa CRNN và CNN khác cả phần conv nên đây không phải phép thử sạch cho tác dụng của GRU.",
    },
    {
        "title": "Performance–cost trade-off",
        "bullets": ["mobilenetv3s: highest clean F1 and fewest MACs (18 M), but slowest on GPU (6.9 ms at batch 1)",
                    "small_cnn: fastest (1.2 ms) and smallest (0.39 M params)",
                    "crnn: dominated on clean F1 for latency and parameters; no single winner",
                    "Peak VRAM was not measured"],
        "image": "rq4_pareto.png",
        "notes": "Câu hỏi E4. Không có người thắng tuyệt đối, tuỳ tiêu chí chi phí. MobileNet có F1 cao nhất và ít MACs nhất nhưng chậm nhất trên GPU ở batch 1. "
                 "small_cnn nhanh nhất và nhỏ nhất. CRNN bị trội trên đường Pareto theo độ trễ và số tham số ở dữ liệu sạch. Chúng em chưa đo VRAM đỉnh.",
    },
    {
        "title": "Proposed Final question (RQ1) and design",
        "bullets": ["How robust are the three models when the clean test set is corrupted at test time?",
                    "Fixed: split, checkpoints (3 seeds), pipeline. No retraining, no test-set tuning",
                    "4 corruptions × 3 severities: pink noise (SNR 20/10/0 dB), gain (−12/−6/+6 dB), time mask (10/25/50 %), Mel-band mask (8/16/32)",
                    "Report F1 clean / corrupted / ΔF1, degradation curves, per-class drops, confusion shifts, flipped clips"],
        "notes": "Câu hỏi cho phần Final là câu RQ1 của đề: ba mô hình mất bao nhiêu khi tập test bị làm bẩn. Giữ nguyên split, checkpoint và pipeline, không huấn luyện lại, "
                 "chỉ thay đổi đầu vào lúc test. Bốn kiểu làm bẩn, mỗi kiểu ba mức. Kết quả sẽ báo cáo F1 sạch, F1 bẩn, ΔF1 kèm đường suy giảm và phân tích theo lớp.",
    },
    {
        "title": "Limitations",
        "bullets": ["Midterm baselines were first run on Colab (T4, one seed); the 3-seed numbers here were run locally (RTX 2060 SUPER)",
                    "Several classes have only 8–17 test clips: one clip moves recall by 6–12 points",
                    "No recording ID: segments from one long recording may be split across partitions",
                    "Listening audit of misclassified clips is not finished"],
        "notes": "Hạn chế: số Colab một seed và số chạy local ba seed không trộn lẫn; nhiều lớp chỉ có 8 đến 17 clip test; "
                 "không có mã phiên ghi âm nên có thể còn rò rỉ nhẹ; phần nghe thử lỗi chưa xong. Em nêu thẳng các điểm này thay vì né.",
    },
]


def add_slide(prs, s):
    layout = prs.slide_layouts[6]
    sl = prs.slides.add_slide(layout)
    W, H = prs.slide_width, prs.slide_height
    tb = sl.shapes.add_textbox(Inches(0.5), Inches(0.3), W - Inches(1.0), Inches(0.9))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.text = s["title"]
    tf.paragraphs[0].runs[0].font.size = Pt(28)
    tf.paragraphs[0].runs[0].font.bold = True
    tf.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)

    has_img = "image" in s
    left_w = Inches(5.6) if has_img else W - Inches(1.0)
    top = Inches(1.4)
    if "table" in s:
        rows, cols = len(s["table"]), len(s["table"][0])
        t_h = Inches(0.4) * rows
        gt = sl.shapes.add_table(rows, cols, Inches(0.5), top, left_w, t_h).table
        for i, row in enumerate(s["table"]):
            for j, v in enumerate(row):
                c = gt.cell(i, j)
                c.text = str(v)
                for p in c.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(12)
        top = top + t_h + Inches(0.2)
    bb = sl.shapes.add_textbox(Inches(0.5), top, left_w, H - top - Inches(0.4))
    bf = bb.text_frame
    bf.word_wrap = True
    for i, b in enumerate(s["bullets"]):
        p = bf.paragraphs[0] if i == 0 else bf.add_paragraph()
        p.text = ("• " if len(s["bullets"]) > 1 or "table" in s else "") + b
        p.font.size = Pt(16 if has_img or "table" in s else 20)
        p.space_after = Pt(8)
    if has_img:
        path = FIG / s["image"]
        if not path.exists():
            sys.exit(f"missing figure {path}")
        max_w, max_h = W - Inches(6.4) - Inches(0.3), H - Inches(1.8)
        pic = sl.shapes.add_picture(str(path), Inches(6.4), Inches(1.4))
        ratio = min(max_w / pic.width, max_h / pic.height)
        pic.width, pic.height = Emu(int(pic.width * ratio)), Emu(int(pic.height * ratio))
    sl.notes_slide.notes_text_frame.text = s["notes"]


def main() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for s in SLIDES:
        add_slide(prs, s)
    OUT_DIR.mkdir(exist_ok=True)
    out = OUT_DIR / "midterm_slides.pptx"
    prs.save(out)

    lines = ["# Kịch bản thuyết trình midterm (tiếng Việt)\n",
             "Slide viết bằng tiếng Anh, lời nói bằng tiếng Việt. Sinh tự động bởi `scripts/build_presentation.py`; "
             "sửa nội dung trong file đó rồi chạy lại. Mục có dấu [CẦN BẠN LÀM] là chỗ tôi không thể làm thay.\n"]
    for i, s in enumerate(SLIDES, 1):
        lines += [f"## Slide {i}: {s['title']}\n", s["notes"] + "\n"]
    SCRIPT_MD.parent.mkdir(exist_ok=True)
    SCRIPT_MD.write_text("\n".join(lines), encoding="utf-8")
    print(f"{len(SLIDES)} slides -> {out}\nscript -> {SCRIPT_MD}")


if __name__ == "__main__":
    main()
